import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from src.excel_writer import write_consolidated_workbook


class ExcelWriterTests(unittest.TestCase):
    def test_writes_workbook_to_output_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "consolidated.xlsx"

            written_path = write_consolidated_workbook(
                output_path,
                pd.DataFrame({"PID": ["P1"]}),
            )

            self.assertEqual(written_path, output_path.resolve())
            workbook = load_workbook(output_path, read_only=True)
            self.assertEqual(workbook.sheetnames, ["Consolidated_Data"])
            self.assertEqual(workbook["Consolidated_Data"]["A2"].value, "P1")
            workbook.close()

    def test_locked_output_has_actionable_error_and_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "consolidated.xlsx"
            output_path.write_bytes(b"existing workbook")

            with patch("src.excel_writer.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaisesRegex(
                    PermissionError,
                    "Close it in Excel or any other application",
                ):
                    write_consolidated_workbook(
                        output_path,
                        pd.DataFrame({"PID": ["P1"]}),
                    )

            self.assertEqual(output_path.read_bytes(), b"existing workbook")
            self.assertEqual(list(Path(temp_dir).iterdir()), [output_path])


if __name__ == "__main__":
    unittest.main()
